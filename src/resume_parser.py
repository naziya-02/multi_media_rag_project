from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path
from typing import Any

import pymupdf as fitz
from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from src.config import MAX_UPLOAD_BYTES
from src.privacy import redact_personal_details
from src.skill_extractor import extract_skills

SECTION_ALIASES = {
    "skills": ("skills", "technical skills", "core competencies", "competencies"),
    "education": ("education", "academic background", "qualifications"),
    "experience": ("experience", "work experience", "employment history", "professional experience"),
    "projects": ("projects", "selected projects", "project experience"),
    "certifications": ("certifications", "certificates", "licenses"),
}
NON_CONTENT_HEADINGS = {
    "about me",
    "awards",
    "contact",
    "interests",
    "languages",
    "objective",
    "profile",
    "professional summary",
    "summary",
}
EMAIL_RE = re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.I)
PHONE_RE = re.compile(r"(?<!\w)(?:\+?\d[\d\s().-]{8,}\d)(?!\w)")
LINK_RE = re.compile(r"(?:https?://)?(?:www\.)?(?:linkedin\.com|github\.com)/\S+", re.I)
class ResumeParseError(ValueError):
    """A clear error raised when the supplied resume cannot be parsed."""


def _extract_text(data: bytes, suffix: str) -> str:
    try:
        if suffix == ".pdf":
            with fitz.open(stream=data, filetype="pdf") as document:
                return "\n".join(page.get_text() for page in document)
        if suffix == ".docx":
            document = Document(io.BytesIO(data))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                paragraphs.extend(cell.text for row in table.rows for cell in row.cells)
            return "\n".join(paragraphs)
    except (fitz.FileDataError, fitz.EmptyFileError) as error:
        raise ResumeParseError("This PDF is damaged, encrypted, or contains no readable pages.") from error
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as error:
        raise ResumeParseError("This DOCX file is invalid or damaged.") from error
    raise ResumeParseError("Unsupported file type. Upload a PDF or DOCX resume.")


def _clean_lines(text: str) -> list[str]:
    return [re.sub(r"\s+", " ", line).strip(" \t•*-|") for line in text.splitlines() if line.strip()]


def _section_for_heading(line: str) -> str | None:
    normalized = re.sub(r"[^a-z ]", "", line.lower()).strip()
    if len(normalized) > 45:
        return None
    for section, headings in SECTION_ALIASES.items():
        if normalized in headings:
            return section
    return None


def _extract_sections(lines: list[str]) -> dict[str, list[str]]:
    sections = {key: [] for key in SECTION_ALIASES}
    current = None
    for line in lines:
        heading = _section_for_heading(line)
        if heading:
            current = heading
            continue
        normalized = re.sub(r"[^a-z ]", "", line.lower()).strip()
        if normalized in NON_CONTENT_HEADINGS:
            current = None
            continue
        if current:
            sections[current].append(line)
    return sections


def parse_resume(data: bytes, filename: str) -> dict[str, Any]:
    suffix = Path(filename).suffix.lower()
    if suffix not in {".pdf", ".docx"}:
        raise ResumeParseError("Unsupported file type. Upload a PDF or DOCX resume.")
    if not data:
        raise ResumeParseError("The uploaded file is empty.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise ResumeParseError("Resume is over the 10 MB upload limit.")
    if suffix == ".pdf" and not data.startswith(b"%PDF-"):
        raise ResumeParseError("The uploaded file does not have a valid PDF signature.")
    if suffix == ".docx" and not data.startswith(b"PK"):
        raise ResumeParseError("The uploaded file does not have a valid DOCX signature.")
    text = _extract_text(data, suffix)
    lines = _clean_lines(text)
    if len(text.strip()) < 20 or not lines:
        raise ResumeParseError("No readable text was found. Scanned-image PDFs need OCR before upload.")

    emails = list(dict.fromkeys(EMAIL_RE.findall(text)))
    phones = list(
        dict.fromkeys(
            match.strip()
            for match in PHONE_RE.findall(text)
            if sum(character.isdigit() for character in match) >= 10
        )
    )
    links = list(dict.fromkeys(match.rstrip(".,)") for match in LINK_RE.findall(text)))
    headings = {heading for values in SECTION_ALIASES.values() for heading in values}
    title_terms = {
        "analyst",
        "architect",
        "consultant",
        "designer",
        "developer",
        "engineer",
        "intern",
        "manager",
        "recruiter",
        "scientist",
        "specialist",
    }
    name = ""
    for line in lines[:8]:
        normalized = re.sub(r"[^a-z ]", "", line.lower()).strip()
        if (
            len(line.split()) in range(2, 5)
            and len(line) < 60
            and normalized not in headings
            and not (set(normalized.split()) & title_terms)
            and not EMAIL_RE.search(line)
            and not PHONE_RE.search(line)
            and not LINK_RE.search(line)
        ):
            name = line
            break

    sections = _extract_sections(lines)
    safe_text = redact_personal_details(text, name)
    skills = list(
        dict.fromkeys(
            extract_skills(redact_personal_details("\n".join(sections["skills"]), name))
            + extract_skills(safe_text)
        )
    )
    return {
        "name": name,
        "email": emails[0] if emails else "",
        "phone": phones[0] if phones else "",
        "links": links,
        "skills": skills,
        "education": sections["education"],
        "experience": sections["experience"],
        "projects": sections["projects"],
        "certifications": sections["certifications"],
        "raw_text": text,
        "filename": Path(filename).name,
    }
